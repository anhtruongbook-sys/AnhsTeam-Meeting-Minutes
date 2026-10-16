# -*- coding: utf-8 -*-
"""
Engine Xuất Bản Biên Bản Cuộc Họp (.docx) Chuẩn Nghị định 30/2020/NĐ-CP
Tuân thủ nghiêm ngặt các quy tắc định dạng hành chính:
1. Khổ giấy A4, lề chuẩn (Trên 20mm, Dưới 20mm, Trái 30mm, Phải 15mm).
2. Phông chữ Times New Roman, màu chữ Đen tuyền (#000000) 100% toàn văn bản và bảng biểu.
3. Tuyệt đối không sử dụng ký tự xuống dòng mềm (\\n) trong các đoạn văn căn đều (JUSTIFY).
4. Quốc hiệu & Tiêu ngữ trình bày bằng bảng 2 cột ẩn viền.
5. Bảng phân công nhiệm vụ nền trắng tự nhiên, viền đen mảnh 0.5pt, lặp lại tiêu đề khi sang trang.
"""

import argparse
import json
import os
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

BLACK = RGBColor(0, 0, 0)
FONT_NAME = "Times New Roman"


def set_run_font(run, size_pt: float, bold: bool = False, italic: bool = False, underline: bool = False):
    """Thiết lập phông chữ Times New Roman và màu đen tuyền (#000000) cho một Run."""
    run.font.name = FONT_NAME
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.underline = underline
    run.font.color.rgb = BLACK
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), FONT_NAME)
    rFonts.set(qn("w:hAnsi"), FONT_NAME)
    rFonts.set(qn("w:cs"), FONT_NAME)
    rFonts.set(qn("w:eastAsia"), FONT_NAME)


def add_clean_paragraph(
    doc,
    text: str = "",
    size_pt: float = 13.0,
    bold: bool = False,
    italic: bool = False,
    alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    space_before_pt: float = 3.0,
    space_after_pt: float = 3.0,
    first_line_indent_cm: float = 0.0,
):
    """
    Tạo một đoạn văn bản độc lập (Hard Return), loại bỏ mọi ký tự xuống dòng mềm (\\n)
    bên trong đoạn căn đều để chống lỗi kéo giãn khoảng cách chữ của Word OpenXML.
    """
    sanitized_text = " ".join(str(text).splitlines()).strip()
    p = doc.add_paragraph()
    p.alignment = alignment
    p.paragraph_format.space_before = Pt(space_before_pt)
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.line_spacing = 1.15
    if first_line_indent_cm > 0:
        p.paragraph_format.first_line_indent = Cm(first_line_indent_cm)
    if sanitized_text:
        run = p.add_run(sanitized_text)
        set_run_font(run, size_pt=size_pt, bold=bold, italic=italic)
    return p


def set_table_borders_invisible(table):
    """Ẩn toàn bộ khung viền cho bảng Quốc hiệu Tiêu ngữ và bảng Chữ ký."""
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement("w:tblBorders")
    for border_name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        border = OxmlElement(f"w:{border_name}")
        border.set(qn("w:val"), "none")
        border.set(qn("w:sz"), "0")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "auto")
        tblBorders.append(border)
    tblPr.append(tblBorders)


def set_table_borders_thin_black(table):
    """Thiết lập khung viền đơn màu đen mảnh 0.5pt (sz=4) cho bảng dữ liệu hành chính."""
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement("w:tblBorders")
    for border_name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        border = OxmlElement(f"w:{border_name}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")  # 4 = 0.5pt
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")
        tblBorders.append(border)
    tblPr.append(tblBorders)


def set_row_repeat_header_and_cant_split(row, is_header: bool = False):
    """Chống cắt ngang dòng khi sang trang và lặp lại dòng tiêu đề bảng."""
    trPr = row._tr.get_or_add_trPr()
    cantSplit = OxmlElement("w:cantSplit")
    trPr.append(cantSplit)
    if is_header:
        tblHeader = OxmlElement("w:tblHeader")
        trPr.append(tblHeader)


def build_meeting_minutes_docx(data: dict, output_docx_path: str):
    """Khởi tạo và xuất tệp Biên bản cuộc họp (.docx) từ dữ liệu cấu trúc JSON."""
    doc = Document()

    # 1. Thiết lập khổ giấy A4 và lề chuẩn Nghị định 30/2020/NĐ-CP
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(30)
    section.right_margin = Mm(15)

    meta = data.get("metadata", {})
    gov_body = meta.get("governing_body", "")
    org_name = meta.get("organization_name", "TÊN CƠ QUAN / ĐƠN VỊ")
    doc_number = meta.get("document_number", "Số: ....../BB-......")
    location_date = meta.get("location_date", "Hà Nội, ngày 16 tháng 10 năm 2026")

    # 2. Bảng Header Quốc hiệu & Tiêu ngữ (2 cột, ẩn viền)
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_invisible(header_table)

    cell_left = header_table.cell(0, 0)
    cell_right = header_table.cell(0, 1)
    cell_left.width = Cm(6.5)
    cell_right.width = Cm(10.0)

    # Cột trái: Tên cơ quan
    cell_left.paragraphs[0].text = ""
    if gov_body:
        p_gov = cell_left.paragraphs[0]
        p_gov.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_gov.paragraph_format.space_before = Pt(0)
        p_gov.paragraph_format.space_after = Pt(2)
        r_gov = p_gov.add_run(gov_body.upper())
        set_run_font(r_gov, size_pt=12.0, bold=False)
        p_org = cell_left.add_paragraph()
    else:
        p_org = cell_left.paragraphs[0]

    p_org.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_org.paragraph_format.space_before = Pt(0)
    p_org.paragraph_format.space_after = Pt(4)
    r_org = p_org.add_run(org_name.upper())
    set_run_font(r_org, size_pt=12.0, bold=True)

    p_num = cell_left.add_paragraph()
    p_num.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_num.paragraph_format.space_before = Pt(2)
    p_num.paragraph_format.space_after = Pt(0)
    r_num = p_num.add_run(doc_number)
    set_run_font(r_num, size_pt=12.0, bold=False)

    # Cột phải: Quốc hiệu & Tiêu ngữ
    p_qh = cell_right.paragraphs[0]
    p_qh.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_qh.paragraph_format.space_before = Pt(0)
    p_qh.paragraph_format.space_after = Pt(2)
    r_qh = p_qh.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
    set_run_font(r_qh, size_pt=12.0, bold=True)

    p_tn = cell_right.add_paragraph()
    p_tn.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tn.paragraph_format.space_before = Pt(0)
    p_tn.paragraph_format.space_after = Pt(4)
    r_tn = p_tn.add_run("Độc lập - Tự do - Hạnh phúc")
    set_run_font(r_tn, size_pt=13.0, bold=True, underline=True)

    p_date = cell_right.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_date.paragraph_format.space_before = Pt(4)
    p_date.paragraph_format.space_after = Pt(0)
    r_date = p_date.add_run(location_date)
    set_run_font(r_date, size_pt=12.5, italic=True)

    # 3. Tên loại văn bản và Trích yếu nội dung
    add_clean_paragraph(
        doc,
        text=meta.get("meeting_title", "BIÊN BẢN CUỘC HỌP"),
        size_pt=14.0,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_before_pt=14.0,
        space_after_pt=2.0,
    )
    add_clean_paragraph(
        doc,
        text=meta.get("meeting_subject", ""),
        size_pt=13.0,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_before_pt=0.0,
        space_after_pt=10.0,
    )

    # 4. Mục I: Thời gian, địa điểm và thành phần tham dự
    add_clean_paragraph(
        doc,
        text="I. THỜI GIAN, ĐỊA ĐIỂM VÀ THÀNH PHẦN THAM DỰ",
        size_pt=13.0,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        space_before_pt=6.0,
        space_after_pt=4.0,
    )

    start_time = meta.get("start_time", "...... giờ ...... phút")
    venue = meta.get("venue", "............................................................")
    add_clean_paragraph(
        doc,
        text=f"1. Thời gian bắt đầu: {start_time}.",
        size_pt=13.0,
        bold=False,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )
    add_clean_paragraph(
        doc,
        text=f"2. Địa điểm / Hình thức họp: {venue}.",
        size_pt=13.0,
        bold=False,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )
    add_clean_paragraph(
        doc,
        text="3. Thành phần tham dự:",
        size_pt=13.0,
        bold=False,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )

    participants = data.get("participants", {})
    chair = participants.get("chairperson", {})
    sec = participants.get("secretary", {})
    attendees = participants.get("attendees", [])

    if chair:
        add_clean_paragraph(
            doc,
            text=f"- Chủ trì: {chair.get('name', '')} - {chair.get('title', '')}.",
            size_pt=13.0,
            alignment=WD_ALIGN_PARAGRAPH.LEFT,
        )
    if sec:
        add_clean_paragraph(
            doc,
            text=f"- Thư ký: {sec.get('name', '')} - {sec.get('title', '')}.",
            size_pt=13.0,
            alignment=WD_ALIGN_PARAGRAPH.LEFT,
        )
    for att in attendees:
        add_clean_paragraph(
            doc,
            text=f"- Thành viên: {att.get('name', '')} - {att.get('title', '')}.",
            size_pt=13.0,
            alignment=WD_ALIGN_PARAGRAPH.LEFT,
        )

    # 5. Mục II: Nội dung cuộc họp (Trình bày theo từng lượt phát biểu)
    add_clean_paragraph(
        doc,
        text="II. NỘI DUNG CUỘC HỌP",
        size_pt=13.0,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        space_before_pt=8.0,
        space_after_pt=4.0,
    )

    presentations = data.get("agenda_and_presentations", [])
    for idx, item in enumerate(presentations, start=1):
        speaker = item.get("speaker", "")
        ts_range = item.get("timestamp_range", "")
        header_line = f"{idx}. Phát biểu của {speaker}"
        if ts_range:
            header_line += f" (Mốc băng ghi âm: {ts_range}):"
        else:
            header_line += ":"

        add_clean_paragraph(
            doc,
            text=header_line,
            size_pt=13.0,
            bold=True,
            alignment=WD_ALIGN_PARAGRAPH.LEFT,
            space_before_pt=4.0,
            space_after_pt=2.0,
        )
        for point in item.get("summary_points", []):
            add_clean_paragraph(
                doc,
                text=f"- {point}",
                size_pt=13.0,
                bold=False,
                alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                first_line_indent_cm=0.75,
            )

    # 6. Mục III: Kết luận cuộc họp và Phân công nhiệm vụ
    add_clean_paragraph(
        doc,
        text="III. KẾT LUẬN CUỘC HỌP VÀ PHÂN CÔNG NHIỆM VỤ",
        size_pt=13.0,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        space_before_pt=8.0,
        space_after_pt=4.0,
    )

    for conc in data.get("conclusions", []):
        add_clean_paragraph(
            doc,
            text=f"- {conc}",
            size_pt=13.0,
            bold=False,
            alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
            first_line_indent_cm=0.75,
        )

    action_items = data.get("action_items", [])
    if action_items:
        add_clean_paragraph(
            doc,
            text="Bảng phân công nhiệm vụ thực hiện:",
            size_pt=13.0,
            bold=True,
            alignment=WD_ALIGN_PARAGRAPH.LEFT,
            space_before_pt=6.0,
            space_after_pt=4.0,
        )

        table = doc.add_table(rows=1, cols=5)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders_thin_black(table)

        headers = ["TT", "Nội dung công việc", "Đơn vị / Cá nhân chủ trì", "Phối hợp", "Thời hạn hoàn thành"]
        col_widths = [Cm(1.2), Cm(6.8), Cm(3.5), Cm(2.5), Cm(2.5)]

        hdr_row = table.rows[0]
        set_row_repeat_header_and_cant_split(hdr_row, is_header=True)
        for col_idx, h_text in enumerate(headers):
            cell = hdr_row.cells[col_idx]
            cell.width = col_widths[col_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(h_text)
            set_run_font(r, size_pt=11.5, bold=True)

        for item in action_items:
            row = table.add_row()
            set_row_repeat_header_and_cant_split(row, is_header=False)
            vals = [
                str(item.get("stt", "")),
                str(item.get("task", "")),
                str(item.get("assignee", "")),
                str(item.get("coordinator", "")),
                str(item.get("deadline", "")),
            ]
            aligns = [
                WD_ALIGN_PARAGRAPH.CENTER,
                WD_ALIGN_PARAGRAPH.JUSTIFY,
                WD_ALIGN_PARAGRAPH.LEFT,
                WD_ALIGN_PARAGRAPH.LEFT,
                WD_ALIGN_PARAGRAPH.CENTER,
            ]
            bolds = [True, False, False, False, True]

            for col_idx, val in enumerate(vals):
                cell = row.cells[col_idx]
                cell.width = col_widths[col_idx]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p = cell.paragraphs[0]
                p.alignment = aligns[col_idx]
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                clean_val = " ".join(val.splitlines()).strip()
                r = p.add_run(clean_val)
                set_run_font(r, size_pt=11.5, bold=bolds[col_idx])

    # 7. Thời gian kết thúc & Bảng Chữ ký xác nhận
    end_time = meta.get("end_time", "...... giờ ...... phút cùng ngày")
    add_clean_paragraph(
        doc,
        text=f"Cuộc họp kết thúc vào lúc {end_time}. Biên bản này đã được đọc lại cho toàn thể thành viên tham dự cùng nghe và thống nhất thông qua./.",
        size_pt=13.0,
        bold=False,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        space_before_pt=8.0,
        space_after_pt=12.0,
        first_line_indent_cm=0.75,
    )

    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_invisible(sig_table)

    sig_left = sig_table.cell(0, 0)
    sig_right = sig_table.cell(0, 1)
    sig_left.width = Cm(8.0)
    sig_right.width = Cm(8.5)

    p_sl_title = sig_left.paragraphs[0]
    p_sl_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sl = p_sl_title.add_run("THƯ KÝ CUỘC HỌP")
    set_run_font(r_sl, size_pt=12.5, bold=True)

    p_sl_sub = sig_left.add_paragraph()
    p_sl_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sl_sub.paragraph_format.space_after = Pt(48)
    r_sl_sub = p_sl_sub.add_run("(Ký và ghi rõ họ tên)")
    set_run_font(r_sl_sub, size_pt=11.5, italic=True)

    p_sl_name = sig_left.add_paragraph()
    p_sl_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sl_name = p_sl_name.add_run(sec.get("name", ""))
    set_run_font(r_sl_name, size_pt=12.5, bold=True)

    p_sr_title = sig_right.paragraphs[0]
    p_sr_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sr = p_sr_title.add_run("CHỦ TRÌ CUỘC HỌP")
    set_run_font(r_sr, size_pt=12.5, bold=True)

    p_sr_sub = sig_right.add_paragraph()
    p_sr_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sr_sub.paragraph_format.space_after = Pt(48)
    r_sr_sub = p_sr_sub.add_run("(Ký, ghi rõ họ tên và đóng dấu nếu có)")
    set_run_font(r_sr_sub, size_pt=11.5, italic=True)

    p_sr_name = sig_right.add_paragraph()
    p_sr_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sr_name = p_sr_name.add_run(chair.get("name", ""))
    set_run_font(r_sr_name, size_pt=12.5, bold=True)

    os.makedirs(os.path.dirname(os.path.abspath(output_docx_path)), exist_ok=True)
    doc.save(output_docx_path)
    print(f"[Hoàn tất] Đã xuất Biên bản cuộc họp chuẩn Nghị định 30/2020/NĐ-CP tại: {output_docx_path}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Xuất Biên bản cuộc họp (.docx) từ file dữ liệu JSON.")
    parser.add_argument("--input-json", required=True, help="Đường dẫn tệp JSON chứa cấu trúc nội dung cuộc họp")
    parser.add_argument("--output-docx", required=True, help="Đường dẫn tệp Word (.docx) đầu ra")
    args = parser.parse_args()

    with open(args.input_json, "r", encoding="utf-8") as f:
        meeting_data = json.load(f)

    build_meeting_minutes_docx(meeting_data, args.output_docx)
